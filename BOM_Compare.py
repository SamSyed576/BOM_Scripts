import os
import re
import sys
import time
import tkinter as tk
from tkinter import filedialog, messagebox
from pathlib import Path
import traceback

try:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import FormulaRule
except ImportError:
    print("ERROR: Required packages are missing.")
    print("Run: py -m pip install pandas openpyxl")
    sys.exit(1)


# ============================================================
# CONFIGURATION
# ============================================================

OUTPUT_FOLDER = r"D:\Princecraft\File Compare\Delta V3-V2"

SUPPORTED_EXTENSIONS = {
    ".xlsx",
    ".xlsm",
    ".csv",
    ".txt",
}


# ============================================================
# COMPARISON KEY PRIORITY
# ============================================================

KEY_COLUMN_CANDIDATES = [
    "ParentAssemblyPath",
    "ChildPath",
    "FullFilePath",
    "Full File Path",
    "FilePath",
    "File Path",
    "Path",
    "FullPath",
    "Full Path",
    "Filename",
    "FileName",
    "File Name",
    "Name",
    "ItemID",
    "Item ID",
    "ItemId",
    "ItemNo",
    "Item No",
    "PartNumber",
    "Part Number",
    "PartNo",
    "Part No",
    "ID",
]


# ============================================================
# EXCEL COLORS
# ============================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

ADDED_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

REMOVED_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6",
)

CHANGED_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

COMMON_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

FILE_A_ONLY_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6",
)

FILE_B_ONLY_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_header(value):
    """
    Normalize a column header for matching.

    Example:
        Parent Assembly Path
        Parent_Assembly_Path
        Parent-Assembly-Path

    all become:
        parentassemblypath
    """

    if value is None:
        return ""

    s = str(value).strip()

    s = s.strip('"').strip("'")

    s = s.lower()

    s = re.sub(
        r"[\s_\-]+",
        "",
        s,
    )

    return s


def build_column_map(df):
    """
    Build:
        normalized header -> actual header
    """

    result = {}

    for col in df.columns:

        normalized = clean_header(col)

        if normalized and normalized not in result:
            result[normalized] = col

    return result


def is_path_column(column_name):
    """
    Determine whether a column represents a file/path value.
    """

    normalized = clean_header(column_name)

    path_terms = [
        "fullfilepath",
        "filepath",
        "fullpath",
        "path",
        "location",
    ]

    return any(
        term in normalized
        for term in path_terms
    )


# ============================================================
# FILE READING
# ============================================================

def detect_delimiter(path):
    """
    Detect delimiter for CSV/TXT files.

    Supported:
        ##
        ,
        |
        tab
        ;
    """

    try:

        with open(
            path,
            "r",
            encoding="utf-8-sig",
            errors="replace",
        ) as f:

            sample = f.read(30000)

    except Exception:

        return ","


    if "##" in sample:

        return "##"


    candidates = [
        ",",
        "|",
        "\t",
        ";",
    ]

    counts = {
        delimiter: sample.count(delimiter)
        for delimiter in candidates
    }

    best = max(
        counts,
        key=counts.get,
    )

    if counts[best] == 0:
        return ","

    return best


def read_data_file(path):

    start = time.perf_counter()

    filename = os.path.basename(path)

    print(
        f"        Reading: {filename}",
        flush=True,
    )

    extension = Path(path).suffix.lower()


    if extension in {
        ".xlsx",
        ".xlsm",
    }:

        df = pd.read_excel(
            path,
            dtype=str,
        )

    else:

        delimiter = detect_delimiter(path)

        df = pd.read_csv(
            path,
            dtype=str,
            sep=delimiter,
            engine="python",
            keep_default_na=False,
            na_filter=False,
            encoding="utf-8-sig",
            on_bad_lines="warn",
        )


    # Clean headers.
    df.columns = [
        str(col).strip()
        for col in df.columns
    ]


    # Replace missing values.
    df = df.fillna("")


    elapsed = time.perf_counter() - start

    print(
        f"        Rows    : {len(df):,}",
        flush=True,
    )

    print(
        f"        Columns : {len(df.columns):,}",
        flush=True,
    )

    print(
        f"        Read time: {elapsed:.1f} sec",
        flush=True,
    )

    return df


# ============================================================
# VECTORISED NORMALIZATION
# ============================================================

def normalize_series(series, path=False):
    """
    Fast vectorised normalization.

    For normal values:
        trim
        remove surrounding quotes
        lowercase

    For paths:
        trim
        remove quotes
        slash -> backslash
        collapse duplicate backslashes
        S:\\R&D-Partage\\ -> Z:\\
        S:\\R&D-Partage -> Z:\\
        remove trailing backslash
    """

    s = series.astype("string")

    s = s.fillna("")

    s = s.str.strip()

    # Remove surrounding double quotes.
    s = s.str.strip('"')

    # Remove surrounding single quotes.
    s = s.str.strip("'")

    s = s.str.lower()


    if path:

        # Convert forward slashes to Windows slashes.
        s = s.str.replace(
            "/",
            "\\",
            regex=False,
        )

        # Collapse repeated backslashes.
        s = s.str.replace(
            r"\\+",
            "\\\\",
            regex=True,
        )

        # ----------------------------------------------------
        # IMPORTANT:
        #
        # Do NOT use:
        #
        #     "z:\\"
        #
        # as a regex replacement.
        #
        # In Python's regex replacement engine, \ is treated
        # as an escape character.
        #
        # Using a lambda returns the literal replacement safely.
        # ----------------------------------------------------

        s = s.str.replace(
            r"^s:\\r&d-partage\\",
            lambda m: "z:\\",
            regex=True,
        )

        # Exact S:\R&D-Partage root.
        s = s.mask(
            s.eq("s:\\r&d-partage"),
            "z:\\",
        )

        # Remove trailing slash except for drive root such as z:\
        s = s.where(
            s.str.len() <= 3,
            s.str.rstrip("\\"),
        )


    return s


def normalize_dataframe_columns(df, column_info):
    """
    Normalize all comparison columns.

    column_info:
        [
            (normalized_name, actual_column_name),
            ...
        ]

    Returns a DataFrame whose columns are normalized names.
    """

    result = {}

    for normalized_name, actual_column in column_info:

        result[normalized_name] = normalize_series(
            df[actual_column],
            path=is_path_column(actual_column),
        )

    if not result:

        return pd.DataFrame(
            index=df.index
        )

    return pd.DataFrame(
        result,
        index=df.index,
    )


# ============================================================
# COMPARISON KEY DETECTION
# ============================================================

def detect_common_key(df_a, df_b):

    map_a = build_column_map(df_a)

    map_b = build_column_map(df_b)


    for candidate in KEY_COLUMN_CANDIDATES:

        normalized_candidate = clean_header(
            candidate
        )

        if (
            normalized_candidate in map_a
            and
            normalized_candidate in map_b
        ):

            return (
                map_a[normalized_candidate],
                map_b[normalized_candidate],
            )


    # --------------------------------------------------------
    # If the priority list failed, show common columns.
    # --------------------------------------------------------

    common = []

    for normalized_name in map_a:

        if normalized_name in map_b:

            common.append(
                (
                    normalized_name,
                    map_a[normalized_name],
                    map_b[normalized_name],
                )
            )


    common_preview = "\n".join(
        f"          {x[1]}  <->  {x[2]}"
        for x in common[:30]
    )


    raise ValueError(
        "\n"
        "No common comparison key could be detected.\n\n"
        "Common columns found:\n"
        f"{common_preview}\n\n"
        "Please add the desired key to "
        "KEY_COLUMN_CANDIDATES."
    )


# ============================================================
# COLUMN ANALYSIS
# ============================================================

def build_column_analysis(df_a, df_b):

    map_a = build_column_map(df_a)

    map_b = build_column_map(df_b)


    rows = []

    seen = set()


    # File A order first.
    for col in df_a.columns:

        normalized = clean_header(col)

        if not normalized:
            continue

        if normalized in seen:
            continue

        seen.add(normalized)


        if normalized in map_b:

            rows.append(
                {
                    "Column": col,
                    "File A Column": col,
                    "File B Column": map_b[normalized],
                    "Status": "Common",
                }
            )

        else:

            rows.append(
                {
                    "Column": col,
                    "File A Column": col,
                    "File B Column": "",
                    "Status": "File A Only",
                }
            )


    # Then File B-only columns.
    for col in df_b.columns:

        normalized = clean_header(col)

        if not normalized:
            continue

        if normalized in seen:
            continue

        seen.add(normalized)


        rows.append(
            {
                "Column": col,
                "File A Column": "",
                "File B Column": col,
                "Status": "File B Only",
            }
        )


    return pd.DataFrame(rows)


# ============================================================
# EMPTY REPORT
# ============================================================

def empty_report(message):

    return pd.DataFrame(
        {
            "Message": [
                message
            ]
        }
    )


# ============================================================
# BUILD UNION DATA
# ============================================================

def build_added_removed_dataframe(
    source_df,
    source_columns,
    other_columns,
    selected_keys,
    source_name,
):
    """
    Build Added or Removed rows without Python loops
    over every row/cell.
    """

    if len(selected_keys) == 0:

        return empty_report(
            f"No files were {source_name.lower()}."
        )


    subset = source_df.loc[
        source_df.index.isin(selected_keys)
    ].copy()


    source_map = build_column_map(
        source_df
    )

    other_map = build_column_map(
        other_columns
    ) if isinstance(other_columns, pd.DataFrame) else {}


    # --------------------------------------------------------
    # Build output in source/union order.
    # --------------------------------------------------------

    output = pd.DataFrame(
        index=subset.index
    )


    seen = set()


    for col in source_columns:

        normalized = clean_header(col)

        if not normalized:
            continue

        if normalized in seen:
            continue

        seen.add(normalized)

        if col in subset.columns:

            output[col] = subset[col].values


    output.reset_index(
        inplace=True
    )

    output.rename(
        columns={
            "__COMPARE_KEY__": "ComparisonKey"
        },
        inplace=True,
    )

    output["ChangeType"] = source_name

    return output


# ============================================================
# MAIN FILE COMPARISON
# ============================================================

def compare_files(
    file_a,
    file_b,
):

    total_start = time.perf_counter()


    print(
        f"      File A: {file_a}",
        flush=True,
    )

    print(
        f"      File B: {file_b}",
        flush=True,
    )


    # ========================================================
    # READ
    # ========================================================

    df_a = read_data_file(
        file_a
    )

    df_b = read_data_file(
        file_b
    )


    # ========================================================
    # DETECT KEY
    # ========================================================

    print(
        "        Detecting comparison key...",
        flush=True,
    )

    key_a, key_b = detect_common_key(
        df_a,
        df_b,
    )

    print(
        f"        Key A: {key_a}",
        flush=True,
    )

    print(
        f"        Key B: {key_b}",
        flush=True,
    )


    key_is_path = (
        is_path_column(key_a)
        or
        is_path_column(key_b)
    )


    # ========================================================
    # NORMALIZE KEYS
    # ========================================================

    print(
        "        Normalizing comparison keys...",
        flush=True,
    )


    key_series_a = normalize_series(
        df_a[key_a],
        path=key_is_path,
    )

    key_series_b = normalize_series(
        df_b[key_b],
        path=key_is_path,
    )


    duplicate_a = int(
        key_series_a.duplicated(
            keep=False
        ).sum()
    )

    duplicate_b = int(
        key_series_b.duplicated(
            keep=False
        ).sum()
    )


    # ========================================================
    # INDEX BY COMPARISON KEY
    # ========================================================

    a = df_a.copy()

    b = df_b.copy()


    a["__COMPARE_KEY__"] = key_series_a

    b["__COMPARE_KEY__"] = key_series_b


    # Keep first row for duplicate keys.
    a = a.drop_duplicates(
        subset="__COMPARE_KEY__",
        keep="first",
    )

    b = b.drop_duplicates(
        subset="__COMPARE_KEY__",
        keep="first",
    )


    a.set_index(
        "__COMPARE_KEY__",
        inplace=True,
    )

    b.set_index(
        "__COMPARE_KEY__",
        inplace=True,
    )


    # ========================================================
    # KEY SETS
    # ========================================================

    keys_a = set(a.index)

    keys_b = set(b.index)


    added_keys = keys_b - keys_a

    removed_keys = keys_a - keys_b

    common_keys = keys_a & keys_b


    print(
        f"        Added keys  : {len(added_keys):,}",
        flush=True,
    )

    print(
        f"        Removed keys: {len(removed_keys):,}",
        flush=True,
    )

    print(
        f"        Common keys : {len(common_keys):,}",
        flush=True,
    )


    # ========================================================
    # COLUMN MAPS
    # ========================================================

    map_a = build_column_map(
        df_a
    )

    map_b = build_column_map(
        df_b
    )


    # ========================================================
    # COLUMN ANALYSIS
    # ========================================================

    column_analysis = build_column_analysis(
        df_a,
        df_b,
    )


    # ========================================================
    # COMMON COMPARISON COLUMNS
    # ========================================================

    comparable_columns = []

    seen = set()


    for col_a in df_a.columns:

        normalized = clean_header(col_a)

        if not normalized:
            continue

        if normalized in seen:
            continue

        seen.add(normalized)


        if (
            normalized in map_b
            and
            col_a != key_a
        ):

            col_b = map_b[normalized]

            comparable_columns.append(
                (
                    normalized,
                    col_a,
                    col_b,
                )
            )


    # ========================================================
    # ADDED
    # ========================================================

    print(
        "        Building Added report...",
        flush=True,
    )


    if added_keys:

        added_df = b.loc[
            b.index.isin(added_keys)
        ].copy()


        output = pd.DataFrame(
            index=added_df.index
        )


        seen_output = set()


        # For Added rows, prefer File B's actual
        # column names.
        for col in df_b.columns:

            normalized = clean_header(col)

            if not normalized:
                continue

            if normalized in seen_output:
                continue

            seen_output.add(normalized)

            output[col] = added_df[col].values


        output.reset_index(
            inplace=True
        )

        output.rename(
            columns={
                "__COMPARE_KEY__":
                    "ComparisonKey"
            },
            inplace=True,
        )

        output["ChangeType"] = "Added"

        added_df = output

    else:

        added_df = empty_report(
            "No rows were added."
        )


    # ========================================================
    # REMOVED
    # ========================================================

    print(
        "        Building Removed report...",
        flush=True,
    )


    if removed_keys:

        removed_df = a.loc[
            a.index.isin(removed_keys)
        ].copy()


        output = pd.DataFrame(
            index=removed_df.index
        )


        seen_output = set()


        # For Removed rows, prefer File A's
        # actual column names.
        for col in df_a.columns:

            normalized = clean_header(col)

            if not normalized:
                continue

            if normalized in seen_output:
                continue

            seen_output.add(normalized)

            output[col] = removed_df[col].values


        output.reset_index(
            inplace=True
        )

        output.rename(
            columns={
                "__COMPARE_KEY__":
                    "ComparisonKey"
            },
            inplace=True,
        )

        output["ChangeType"] = "Removed"

        removed_df = output

    else:

        removed_df = empty_report(
            "No rows were removed."
        )


    # ========================================================
    # COMMON ROW COMPARISON
    # ========================================================

    print(
        "        Comparing common rows...",
        flush=True,
    )


    if not common_keys:

        changed_df = empty_report(
            "No common keys were found."
        )

    elif not comparable_columns:

        changed_df = empty_report(
            "No common columns available for comparison."
        )

    else:

        # ----------------------------------------------------
        # Preserve index order from File A.
        # ----------------------------------------------------

        common_index = a.index[
            a.index.isin(common_keys)
        ]


        # ----------------------------------------------------
        # Build normalized comparison DataFrames.
        #
        # This replaces the old:
        #
        #   row -> column -> .at[]
        #
        # nested loop.
        #
        # The comparison is now vectorized.
        # ----------------------------------------------------

        normalized_a = {}

        normalized_b = {}


        for normalized_name, col_a, col_b in comparable_columns:

            normalized_a[normalized_name] = normalize_series(
                a.loc[
                    common_index,
                    col_a,
                ],
                path=is_path_column(col_a),
            )

            normalized_b[normalized_name] = normalize_series(
                b.loc[
                    common_index,
                    col_b,
                ],
                path=is_path_column(col_b),
            )


        norm_a = pd.DataFrame(
            normalized_a,
            index=common_index,
        )

        norm_b = pd.DataFrame(
            normalized_b,
            index=common_index,
        )


        # ----------------------------------------------------
        # Vectorized difference mask.
        # ----------------------------------------------------

        diff_mask = norm_a.ne(
            norm_b
        )


        changed_row_mask = diff_mask.any(
            axis=1
        )


        changed_indexes = common_index[
            changed_row_mask
        ]


        unchanged_count = int(
            (~changed_row_mask).sum()
        )


        changed_count = len(
            changed_indexes
        )


        print(
            f"        Changed rows  : {changed_count:,}",
            flush=True,
        )

        print(
            f"        Unchanged rows: {unchanged_count:,}",
            flush=True,
        )


        # ----------------------------------------------------
        # Build Changed report.
        # ----------------------------------------------------

        if changed_count == 0:

            changed_df = empty_report(
                "No changed rows were found."
            )

        else:

            changed_df = pd.DataFrame(
                index=changed_indexes
            )


            # Comparison key.
            changed_df[
                "ComparisonKey"
            ] = changed_indexes


            # ------------------------------------------------
            # Changed columns list.
            #
            # Only one Python iteration per row rather than
            # one Python operation per row x column.
            # ------------------------------------------------

            diff_array = diff_mask.loc[
                changed_indexes
            ].to_numpy(
                dtype=bool
            )


            comparison_column_names = [
                col_a
                for (
                    normalized_name,
                    col_a,
                    col_b,
                )
                in comparable_columns
            ]


            changed_column_names = []


            for row_changes in diff_array:

                names = [
                    comparison_column_names[i]
                    for i, changed in enumerate(
                        row_changes
                    )
                    if changed
                ]

                changed_column_names.append(
                    ", ".join(names)
                )


            changed_df[
                "ChangedColumns"
            ] = changed_column_names


            # ------------------------------------------------
            # Add A/B values only for columns that changed.
            # ------------------------------------------------

            used_output_names = set()


            for (
                normalized_name,
                col_a,
                col_b,
            ) in comparable_columns:

                column_changed = diff_mask.loc[
                    changed_indexes,
                    normalized_name,
                ].to_numpy(
                    dtype=bool
                )


                if not column_changed.any():
                    continue


                output_a_name = (
                    f"{col_a} (File A)"
                )

                output_b_name = (
                    f"{col_b} (File B)"
                )


                # Avoid duplicate output names.
                if output_a_name in used_output_names:
                    suffix = 2

                    base = output_a_name

                    while (
                        f"{base} [{suffix}]"
                        in used_output_names
                    ):
                        suffix += 1

                    output_a_name = (
                        f"{base} [{suffix}]"
                    )


                if output_b_name in used_output_names:
                    suffix = 2

                    base = output_b_name

                    while (
                        f"{base} [{suffix}]"
                        in used_output_names
                    ):
                        suffix += 1

                    output_b_name = (
                        f"{base} [{suffix}]"
                    )


                used_output_names.add(
                    output_a_name
                )

                used_output_names.add(
                    output_b_name
                )


                values_a = (
                    a.loc[
                        changed_indexes,
                        col_a,
                    ]
                    .astype("string")
                    .fillna("")
                    .to_numpy()
                )

                values_b = (
                    b.loc[
                        changed_indexes,
                        col_b,
                    ]
                    .astype("string")
                    .fillna("")
                    .to_numpy()
                )


                # Only show values when that particular
                # column changed for that row.
                values_a = [
                    values_a[i]
                    if column_changed[i]
                    else ""
                    for i in range(
                        len(column_changed)
                    )
                ]

                values_b = [
                    values_b[i]
                    if column_changed[i]
                    else ""
                    for i in range(
                        len(column_changed)
                    )
                ]


                changed_df[
                    output_a_name
                ] = values_a

                changed_df[
                    output_b_name
                ] = values_b


            changed_df[
                "ChangeType"
            ] = "Changed"


            changed_df.reset_index(
                drop=True,
                inplace=True,
            )


    # ========================================================
    # SUMMARY
    # ========================================================

    elapsed = time.perf_counter() - total_start


    summary_df = pd.DataFrame(
        [
            {
                "File A": os.path.basename(file_a),
                "File B": os.path.basename(file_b),
                "Key A": key_a,
                "Key B": key_b,
                "Key Type": (
                    "Path"
                    if key_is_path
                    else "Value"
                ),
                "File A Rows": len(df_a),
                "File B Rows": len(df_b),
                "File A Columns": len(df_a.columns),
                "File B Columns": len(df_b.columns),
                "Common Columns": int(
                    (
                        column_analysis[
                            "Status"
                        ] == "Common"
                    ).sum()
                ),
                "File A Only Columns": int(
                    (
                        column_analysis[
                            "Status"
                        ] == "File A Only"
                    ).sum()
                ),
                "File B Only Columns": int(
                    (
                        column_analysis[
                            "Status"
                        ] == "File B Only"
                    ).sum()
                ),
                "Added Rows": len(
                    added_keys
                ),
                "Removed Rows": len(
                    removed_keys
                ),
                "Changed Rows": (
                    changed_count
                    if (
                        "changed_count"
                        in locals()
                    )
                    else 0
                ),
                "Unchanged Rows": (
                    unchanged_count
                    if (
                        "unchanged_count"
                        in locals()
                    )
                    else 0
                ),
                "Duplicate Key Rows - File A":
                    duplicate_a,
                "Duplicate Key Rows - File B":
                    duplicate_b,
                "Comparison Time (sec)":
                    round(elapsed, 1),
            }
        ]
    )


    print(
        f"        Comparison calculation complete "
        f"({elapsed:.1f} seconds)",
        flush=True,
    )


    return {
        "Summary": summary_df,
        "Column Analysis": column_analysis,
        "Added": added_df,
        "Removed": removed_df,
        "Changed": changed_df,
    }


# ============================================================
# OUTPUT FILE NAME
# ============================================================

def get_unique_output_path(
    output_folder,
    base_name,
):

    os.makedirs(
        output_folder,
        exist_ok=True,
    )


    path = os.path.join(
        output_folder,
        base_name,
    )


    if not os.path.exists(path):

        return path


    stem = Path(base_name).stem

    suffix = Path(base_name).suffix


    counter = 2


    while True:

        candidate = os.path.join(
            output_folder,
            f"{stem}_{counter}{suffix}",
        )

        if not os.path.exists(candidate):

            return candidate

        counter += 1


# ============================================================
# EXCEL WRITING
# ============================================================

def write_dataframes_to_excel(
    reports,
    output_path,
):

    print(
        "        Writing Excel workbook...",
        flush=True,
    )


    # --------------------------------------------------------
    # ONLY FIVE SHEETS.
    #
    # Unchanged has intentionally been removed.
    # --------------------------------------------------------

    sheet_order = [
        "Summary",
        "Column Analysis",
        "Added",
        "Removed",
        "Changed",
    ]


    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:

        for sheet_name in sheet_order:

            reports[
                sheet_name
            ].to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )


    format_report(
        output_path
    )


# ============================================================
# FAST EXCEL FORMATTING
# ============================================================

def find_column_number(
    ws,
    column_name,
):

    for cell in ws[1]:

        if (
            str(cell.value).strip()
            == column_name
        ):

            return cell.column


    return None


def add_status_conditional_formatting(
    ws,
    status_column_name,
    status_values,
):

    if ws.max_row < 2:
        return


    status_column = find_column_number(
        ws,
        status_column_name,
    )


    if status_column is None:
        return


    status_letter = get_column_letter(
        status_column
    )

    last_letter = get_column_letter(
        ws.max_column
    )


    data_range = (
        f"A2:{last_letter}{ws.max_row}"
    )


    for status_value, fill in status_values:

        formula = (
            f'${status_letter}2="{status_value}"'
        )


        ws.conditional_formatting.add(
            data_range,
            FormulaRule(
                formula=[formula],
                fill=fill,
            ),
        )


def calculate_column_widths(
    ws,
    sample_rows=100,
):

    """
    Calculate reasonable widths without scanning
    every cell in very large worksheets.

    This deliberately samples only the first 100 rows.
    """

    max_sample_row = min(
        ws.max_row,
        sample_rows + 1,
    )


    for column_cells in ws.iter_cols(
        min_row=1,
        max_row=max_sample_row,
    ):

        if not column_cells:
            continue


        column_index = (
            column_cells[0].column
        )


        header = column_cells[0].value

        header_text = (
            str(header)
            if header is not None
            else ""
        )


        max_length = len(
            header_text
        )


        for cell in column_cells[1:]:

            value = cell.value

            if value is None:
                continue

            text = str(value)


            # Avoid extremely expensive width calculations
            # for huge values.
            if len(text) > 100:
                text = text[:100]


            if len(text) > max_length:

                max_length = len(text)


        width = min(
            max(
                max_length + 2,
                10,
            ),
            75,
        )


        if is_path_column(
            header_text
        ):

            width = max(
                width,
                35,
            )


        ws.column_dimensions[
            get_column_letter(column_index)
        ].width = width


def format_report(
    output_path,
):

    print(
        "        Formatting Excel workbook...",
        flush=True,
    )


    wb = load_workbook(
        output_path
    )


    # ========================================================
    # SUMMARY
    # ========================================================

    if "Summary" in wb.sheetnames:

        ws = wb["Summary"]

        print(
            "          Formatting sheet: Summary",
            flush=True,
        )


        ws.freeze_panes = "A2"

        ws.auto_filter.ref = ws.dimensions


        for cell in ws[1]:

            cell.fill = HEADER_FILL

            cell.font = HEADER_FONT

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


        # Summary is small, so body alignment is fine.
        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


        calculate_column_widths(
            ws,
            sample_rows=100,
        )


    # ========================================================
    # COLUMN ANALYSIS
    # ========================================================

    if "Column Analysis" in wb.sheetnames:

        ws = wb["Column Analysis"]

        print(
            "          Formatting sheet: Column Analysis",
            flush=True,
        )


        ws.freeze_panes = "A2"

        ws.auto_filter.ref = ws.dimensions


        for cell in ws[1]:

            cell.fill = HEADER_FILL

            cell.font = HEADER_FONT

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


        add_status_conditional_formatting(
            ws,
            "Status",
            [
                (
                    "Common",
                    COMMON_FILL,
                ),
                (
                    "File A Only",
                    FILE_A_ONLY_FILL,
                ),
                (
                    "File B Only",
                    FILE_B_ONLY_FILL,
                ),
            ],
        )


        calculate_column_widths(
            ws,
            sample_rows=100,
        )


    # ========================================================
    # DATA SHEETS
    # ========================================================

    for sheet_name, fill in [
        ("Added", ADDED_FILL),
        ("Removed", REMOVED_FILL),
        ("Changed", CHANGED_FILL),
    ]:

        if sheet_name not in wb.sheetnames:
            continue


        ws = wb[sheet_name]


        print(
            f"          Formatting sheet: {sheet_name}",
            flush=True,
        )


        ws.freeze_panes = "A2"


        if ws.max_row >= 1:

            ws.auto_filter.ref = ws.dimensions


        # ----------------------------------------------------
        # Header only.
        #
        # We intentionally DO NOT loop through every body
        # cell. This is what caused the previous version to
        # spend many minutes formatting huge worksheets.
        # ----------------------------------------------------

        for cell in ws[1]:

            cell.fill = HEADER_FILL

            cell.font = HEADER_FONT

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


        # ----------------------------------------------------
        # Use conditional formatting for the entire data range.
        #
        # This replaces potentially millions of individual
        # cell formatting operations.
        # ----------------------------------------------------

        if sheet_name in {
            "Added",
            "Removed",
            "Changed",
        }:

            add_status_conditional_formatting(
                ws,
                "ChangeType",
                [
                    (
                        sheet_name,
                        fill,
                    ),
                ],
            )


        calculate_column_widths(
            ws,
            sample_rows=100,
        )


    # ========================================================
    # SAVE
    # ========================================================

    wb.save(
        output_path
    )

    wb.close()


# ============================================================
# FOLDER SCANNING
# ============================================================

def get_folder_files(folder):

    files = {}


    for entry in os.scandir(folder):

        if not entry.is_file():
            continue


        extension = (
            Path(entry.name)
            .suffix
            .lower()
        )


        if extension not in SUPPORTED_EXTENSIONS:
            continue


        normalized_name = (
            entry.name.lower()
        )


        if normalized_name not in files:

            files[
                normalized_name
            ] = entry.path


    return files


# ============================================================
# MASTER SUMMARY
# ============================================================

def create_master_summary(
    master_rows,
    folder_a,
    folder_b,
):

    print(
        "Creating master comparison summary...",
        flush=True,
    )


    if master_rows:

        summary_df = pd.DataFrame(
            master_rows
        )

    else:

        summary_df = pd.DataFrame(
            {
                "Message": [
                    "No common files were found."
                ]
            }
        )


    output_path = get_unique_output_path(
        OUTPUT_FOLDER,
        "Folder_Comparison_Summary.xlsx",
    )


    reports = {
        "Summary": summary_df,
    }


    # --------------------------------------------------------
    # The master workbook contains only one summary sheet.
    # --------------------------------------------------------

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:

        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )


    # --------------------------------------------------------
    # Lightweight formatting.
    # --------------------------------------------------------

    wb = load_workbook(
        output_path
    )


    ws = wb["Summary"]


    ws.freeze_panes = "A2"

    ws.auto_filter.ref = ws.dimensions


    for cell in ws[1]:

        cell.fill = HEADER_FILL

        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )


    calculate_column_widths(
        ws,
        sample_rows=200,
    )


    wb.save(
        output_path
    )

    wb.close()


    print(
        f"Master summary created:\n"
        f"  {output_path}",
        flush=True,
    )


    return output_path


# ============================================================
# GUI
# ============================================================

def choose_folder(
    title,
):

    return filedialog.askdirectory(
        title=title
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 70)
    print("BOM FOLDER COMPARATOR")
    print("=" * 70)
    print()


    # --------------------------------------------------------
    # Create Tkinter root.
    # --------------------------------------------------------

    root = tk.Tk()

    root.withdraw()


    try:

        # ====================================================
        # FOLDER A
        # ====================================================

        folder_a = choose_folder(
            "Select Folder A"
        )


        if not folder_a:

            print(
                "Folder A selection cancelled."
            )

            return


        # ====================================================
        # FOLDER B
        # ====================================================

        folder_b = choose_folder(
            "Select Folder B"
        )


        if not folder_b:

            print(
                "Folder B selection cancelled."
            )

            return


        # ====================================================
        # VALIDATE FOLDERS
        # ====================================================

        if os.path.abspath(
            folder_a
        ).lower() == os.path.abspath(
            folder_b
        ).lower():

            messagebox.showerror(
                "Invalid Selection",
                "Folder A and Folder B must be different.",
            )

            return


        print(
            "Folder A:"
        )

        print(
            f"  {folder_a}"
        )

        print()


        print(
            "Folder B:"
        )

        print(
            f"  {folder_b}"
        )

        print()


        print(
            "Output:"
        )

        print(
            f"  {OUTPUT_FOLDER}"
        )

        print()


        os.makedirs(
            OUTPUT_FOLDER,
            exist_ok=True,
        )


        # ====================================================
        # SCAN FOLDER A
        # ====================================================

        print(
            "Scanning Folder A...",
            flush=True,
        )


        files_a = get_folder_files(
            folder_a
        )


        print(
            f"  Files found: {len(files_a)}",
            flush=True,
        )


        # ====================================================
        # SCAN FOLDER B
        # ====================================================

        print(
            "Scanning Folder B...",
            flush=True,
        )


        files_b = get_folder_files(
            folder_b
        )


        print(
            f"  Files found: {len(files_b)}",
            flush=True,
        )


        # ====================================================
        # MATCHING
        # ====================================================

        common_names = sorted(
            set(files_a)
            &
            set(files_b)
        )


        only_a = sorted(
            set(files_a)
            -
            set(files_b)
        )


        only_b = sorted(
            set(files_b)
            -
            set(files_a)
        )


        print()
        print("=" * 70)
        print("FOLDER MATCHING")
        print("=" * 70)


        print(
            f"Common files : {len(common_names)}"
        )

        print(
            f"Folder A only: {len(only_a)}"
        )

        print(
            f"Folder B only: {len(only_b)}"
        )


        # ====================================================
        # MASTER SUMMARY ROWS
        # ====================================================

        master_rows = []


        # ====================================================
        # COMPARE COMMON FILES
        # ====================================================

        print()
        print("=" * 70)
        print(
            f"COMPARING {len(common_names)} COMMON FILES"
        )
        print("=" * 70)


        for number, normalized_name in enumerate(
            common_names,
            start=1,
        ):

            filename = os.path.basename(
                files_a[normalized_name]
            )


            print()
            print("=" * 70)

            print(
                f"[{number}/{len(common_names)}] "
                f"Comparing: {filename}"
            )

            print("=" * 70)


            start_file = time.perf_counter()


            try:

                root.title(
                    f"BOM Comparator - "
                    f"{number}/{len(common_names)} - "
                    f"{filename}"
                )


                root.update()


                # ------------------------------------------------
                # Compare.
                # ------------------------------------------------

                reports = compare_files(
                    files_a[
                        normalized_name
                    ],
                    files_b[
                        normalized_name
                    ],
                )


                # ------------------------------------------------
                # Create output workbook.
                # ------------------------------------------------

                stem = Path(
                    filename
                ).stem


                output_filename = (
                    f"{stem}_Comparison.xlsx"
                )


                output_path = (
                    get_unique_output_path(
                        OUTPUT_FOLDER,
                        output_filename,
                    )
                )


                print(
                    "        Creating comparison workbook...",
                    flush=True,
                )


                write_dataframes_to_excel(
                    reports,
                    output_path,
                )


                elapsed_file = (
                    time.perf_counter()
                    -
                    start_file
                )


                # ------------------------------------------------
                # Extract summary.
                # ------------------------------------------------

                summary = (
                    reports[
                        "Summary"
                    ].iloc[0].to_dict()
                )


                master_row = {
                    "File Name": filename,
                    "File A Rows": summary.get(
                        "File A Rows",
                        "",
                    ),
                    "File B Rows": summary.get(
                        "File B Rows",
                        "",
                    ),
                    "File A Columns": summary.get(
                        "File A Columns",
                        "",
                    ),
                    "File B Columns": summary.get(
                        "File B Columns",
                        "",
                    ),
                    "Common Columns": summary.get(
                        "Common Columns",
                        "",
                    ),
                    "File A Only Columns": summary.get(
                        "File A Only Columns",
                        "",
                    ),
                    "File B Only Columns": summary.get(
                        "File B Only Columns",
                        "",
                    ),
                    "Added Rows": summary.get(
                        "Added Rows",
                        "",
                    ),
                    "Removed Rows": summary.get(
                        "Removed Rows",
                        "",
                    ),
                    "Changed Rows": summary.get(
                        "Changed Rows",
                        "",
                    ),
                    "Unchanged Rows": summary.get(
                        "Unchanged Rows",
                        "",
                    ),
                    "Duplicate Key Rows - File A":
                        summary.get(
                            "Duplicate Key Rows - File A",
                            "",
                        ),
                    "Duplicate Key Rows - File B":
                        summary.get(
                            "Duplicate Key Rows - File B",
                            "",
                        ),
                    "Comparison Time (sec)": round(
                        elapsed_file,
                        1,
                    ),
                    "Output Workbook": output_path,
                }


                master_rows.append(
                    master_row
                )


                print()
                print(
                    f"        Output created:"
                )

                print(
                    f"        {output_path}"
                )

                print()

                print(
                    f"        Added   : "
                    f"{summary.get('Added Rows', 0):,}"
                )

                print(
                    f"        Removed : "
                    f"{summary.get('Removed Rows', 0):,}"
                )

                print(
                    f"        Changed : "
                    f"{summary.get('Changed Rows', 0):,}"
                )

                print(
                    f"        Same    : "
                    f"{summary.get('Unchanged Rows', 0):,}"
                )

                print(
                    f"        Total time: "
                    f"{elapsed_file:.1f} sec"
                )


            except Exception as exc:

                print()
                print(
                    "      ERROR:"
                )

                print(
                    f"      {type(exc).__name__}: "
                    f"{exc}"
                )


                traceback.print_exc()


                master_rows.append(
                    {
                        "File Name": filename,
                        "Status": "ERROR",
                        "Error Type":
                            type(exc).__name__,
                        "Error Message":
                            str(exc),
                    }
                )


            root.update()


        # ====================================================
        # MASTER SUMMARY
        # ====================================================

        print()
        print(
            "Creating master comparison summary...",
            flush=True,
        )


        master_path = create_master_summary(
            master_rows,
            folder_a,
            folder_b,
        )


        print()
        print("=" * 70)
        print("COMPARISON COMPLETE")
        print("=" * 70)


        print()
        print(
            f"Individual reports:"
        )

        print(
            f"  {OUTPUT_FOLDER}"
        )


        print()
        print(
            f"Master summary:"
        )

        print(
            f"  {master_path}"
        )


        print()


        messagebox.showinfo(
            "BOM Comparison Complete",
            (
                f"Comparison complete.\n\n"
                f"Common files: {len(common_names)}\n"
                f"Folder A only: {len(only_a)}\n"
                f"Folder B only: {len(only_b)}\n\n"
                f"Reports saved to:\n"
                f"{OUTPUT_FOLDER}"
            ),
        )


    except Exception as exc:

        print()
        print("=" * 70)
        print("FATAL ERROR")
        print("=" * 70)

        print(
            f"{type(exc).__name__}: {exc}"
        )

        traceback.print_exc()


        try:

            messagebox.showerror(
                "BOM Comparator Error",
                (
                    f"{type(exc).__name__}: "
                    f"{exc}"
                ),
            )

        except Exception:
            pass


    finally:

        try:
            root.destroy()
        except Exception:
            pass


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    main()